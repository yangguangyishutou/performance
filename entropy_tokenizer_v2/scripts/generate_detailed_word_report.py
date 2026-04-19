from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
OUT_PATH = DOCS_DIR / "etv2_detailed_conclusion_report_2026-04-19.docx"


def _set_run_font(run, *, name: str = "Microsoft YaHei", size: float = 10.5, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def _set_doc_defaults(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_run_font(run, size=18, bold=True)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    size = {1: 15, 2: 13, 3: 11}.get(level, 10.5)
    _set_run_font(run, size=size, bold=True)
    p.style = f"Heading {min(level, 3)}"


def add_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        _set_run_font(run)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        _set_run_font(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        p = hdr[idx].paragraphs[0]
        run = p.add_run(header)
        _set_run_font(run, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            run = p.add_run(value)
            _set_run_font(run)


def add_code_block(doc: Document, text: str) -> None:
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        run = p.add_run(line)
        _set_run_font(run, name="Consolas", size=9.5)


def _build_report(doc: Document) -> None:
    add_title(doc, "Entropy Tokenizer V2 详细实验结论报告（可读详版）")
    add_paragraph(doc, "汇报日期：2026-04-19")
    add_paragraph(doc, "实验环境：SeetaCloud 4090 远端；Qwen/Qwen2.5-Coder-1.5B；4bit QLoRA")
    add_paragraph(doc, "本版目标：不仅给结论，还要把每个配置、每条规则、每个优劣点、每类样例解释清楚。")

    add_heading(doc, "一、先给最终结论（防止看完还抓不住重点）", level=1)
    add_bullets(
        doc,
        [
            "最佳训练变体：CPT(77 rows) -> SFT(220 rows)。",
            "首次拿到“端到端真实 prompt token 下降”：context -12.4%，生成期 prompt -10.9%。",
            "准确率仍有损失：raw pass@1 = 0.5，compressed pass@1 = 0.4（-0.1 绝对差距）。",
            "因此当前不是“失败”，也不是“已完成”；更准确是“已跨过成本门槛，未跨过能力门槛”。",
        ],
    )

    add_heading(doc, "二、你现在看到的配置名到底是什么意思", level=1)
    add_paragraph(
        doc,
        "我们这套系统里最容易看不懂的是配置名字。本节把常见名字拆成“中文含义 + 影响范围 + 好处 + 代价”。"
    )
    add_table(
        doc,
        headers=["配置项", "值", "含义", "好处", "代价/风险"],
        rows=[
            ["stage3_backend", "hybrid_ab", "Stage3 使用 A/B 双通道框架", "可控、可解释，可扩展", "配置项更多，调参复杂"],
            ["ET_STAGE3_AB_MODE", "exact_only", "只开 A 通道 exact alias", "稳定，误伤低，便于回滚", "free-text 不压，天花板略低"],
            ["ET_STAGE3_AB_MODE", "hybrid", "同时开 A+B 通道", "理论压缩上限更高", "收益不稳定，行为更难控"],
            ["STAGE2_HYBRID_AB_PROFILE", "stage2_hybrid_ab_aggressive", "Stage2 用激进清洗配置", "静态压缩率显著提升", "可读性下降，语义风险需 guardrail"],
            ["STAGE2_HYBRID_AB_MODE", "blockwise", "按代码块而非逐行清洗", "规则一致性更好", "细粒度可控性稍弱"],
            ["a_min_occ", "3", "A 通道最小出现次数阈值", "避免低频词引入无意义别名", "可能漏压某些短样本"],
            ["a_cost_mode", "context_aware", "按上下文真实成本决策", "不容易被表面频次欺骗", "计算更复杂"],
            ["enable_global_guardrail", "true", "文件级防回退", "防止压后反而变长", "少量额外计算开销"],
            ["enable_incremental_rollback", "true", "逐项回滚别名", "坏 alias 可局部撤销", "实现链路更复杂"],
        ],
    )

    add_heading(doc, "三、规则到底在做什么（通俗解释）", level=1)
    add_heading(doc, "3.1 Stage1：语法骨架压缩（不是变量名压缩）", level=2)
    add_paragraph(
        doc,
        "Stage1 把高频语句头抽象成骨架模板。例如 def/if/for/return 这种语句头会变成 <SYN_x>。它的目标是先压重复结构，再交给后续阶段压细节。"
    )
    add_table(
        doc,
        headers=["Top 骨架", "频次", "有效净节省", "直觉解释"],
        rows=[
            ["def {0}({1}, {2}):", "112", "318", "函数签名重复非常多，最值钱"],
            ["def {0}({1}):", "155", "280", "单参数函数也大量重复"],
            ["if {0} in {1}:", "62", "109", "成员判断模式高频"],
            ["if {0}({1}) > 0:", "61", "104", "轻量判定模板复用高"],
            ["for {0}, {1} in {2}({3}).items():", "59", "95", "字典迭代结构频繁出现"],
        ],
    )

    add_heading(doc, "3.2 Stage2：清洗规则（R01-R05）", level=2)
    add_table(
        doc,
        headers=["规则", "作用", "例子", "保留/删除原则"],
        rows=[
            ["R01", "注释清洗", "# explain x 可删，# noqa 保留", "功能指令注释保留，解释性注释可删"],
            ["R02", "删除空行", "多个空行压成紧凑序列", "默认可删"],
            ["R03", "删除尾随空白", "行末多余空格删除", "默认可删"],
            ["R04", "去缩进（激进）", "把可视缩进压扁", "用于压缩评测，不是源码保真路径"],
            ["R05", "低风险 docstring 删除", "私有 helper docstring 可删", "只删低风险，保守判定"],
        ],
    )
    add_paragraph(
        doc,
        "一句话记忆：Stage2 的本质是“去掉低价值布局与注释负担”，但不会无脑删；指令注释（如 noqa、pragma）会被保留。"
    )

    add_heading(doc, "3.3 Stage3 A/B 路由：为什么有些东西压，有些不压", level=2)
    add_table(
        doc,
        headers=["对象类型", "默认路由", "为什么"],
        rows=[
            ["变量名/属性名", "A", "字面别名最安全，收益稳定"],
            ["identifier-like 字符串", "A", "更像代码 token，适合 exact alias"],
            ["path/url/regex-like", "A", "结构化文本，误伤可控"],
            ["free-text 句子", "B", "需要句子级处理，exact 风险高"],
            ["未知高风险", "不压/回滚", "宁可不压也不引入错误"],
        ],
    )
    add_paragraph(
        doc,
        "你看到“为什么这个字符串没压”通常就落在这里：不是系统没看到，而是被路由到 B；若当前是 exact_only，B 关闭，就会故意不压。"
    )

    add_heading(doc, "四、真实样例 A：主线压缩前后到底发生了什么", level=1)
    add_paragraph(doc, "样例代码（原始）：")
    add_code_block(
        doc,
        """
def hydrate_customer_identifier(customer_identifier_payload, customer_identifier_lookup, customer_identifier_cache):
    customer_identifier_payload = customer_identifier_lookup.get(customer_identifier_payload, customer_identifier_payload)
    customer_identifier_cache[customer_identifier_payload] = customer_identifier_payload
    customer_identifier_lookup[customer_identifier_payload] = customer_identifier_payload
    customer_identifier_payload = customer_identifier_payload.strip()
    audit_event_label = "customer identifier payload mismatch"
    failure_message = "customer identifier payload mismatch"
    fallback_message = "customer identifier payload mismatch"
    return customer_identifier_payload
        """,
    )
    add_table(
        doc,
        headers=["阶段", "tokens", "变化说明"],
        rows=[
            ["原始", "109", "未压缩"],
            ["Stage1", "102", "函数签名和 return 等语法骨架化"],
            ["Stage2", "98", "布局/空白进一步收紧"],
            ["Stage3", "71", "A 通道别名引入，得到主要收益"],
        ],
    )
    add_paragraph(doc, "最终 Stage3 序列：")
    add_code_block(
        doc,
        """
<SYN_11> hydrate_customer_identifier a b customer_identifier_cache
a = b.get(a, a)
customer_identifier_cache[a] = a
b[a] = a
a = a.strip()
<c> audit_event_label "customer identifier payload mismatch"
<c> failure_message "customer identifier payload mismatch"
<c> fallback_message "customer identifier payload mismatch"
<SYN_19> a
        """,
    )
    add_bullets(
        doc,
        [
            "customer_identifier_payload 被压成 a（高频，净收益为正）。",
            "customer_identifier_lookup 被压成 b（达到 min_occ，收益略正）。",
            "customer_identifier_cache 没压：出现次数不足（低于 a_min_occ=3）。",
            "三条报错 free-text 没压：被路由为 B 候选，而 exact_only 默认关 B。",
        ],
    )

    add_heading(doc, "五、真实样例 B：Stage2 规则怎么保留“该保留的东西”", level=1)
    add_paragraph(doc, "原始片段：")
    add_code_block(
        doc,
        """
#!/usr/bin/env python
# noqa: F401
from customer_service import fetch_records

def _build_customer_map(customer_identifier, raw_records):
    \"\"\"private helper docstring for removable setup\"\"\"
    # remove this explanatory comment
    normalized_rows = {}
    for record_key, record_value in fetch_records(raw_records).items():
        if record_key in raw_records:
            normalized_rows[customer_identifier] = record_value
    return normalized_rows
        """,
    )
    add_bullets(
        doc,
        [
            "shebang 和 # noqa 会保留（指令注释）。",
            "普通解释性注释会删除（R01）。",
            "低风险私有 docstring 会删除（R05）。",
            "布局空白与缩进会按模式压缩（R02/R03/R04）。",
        ],
    )
    add_paragraph(
        doc,
        "这就是“有规则地删”，不是“无差别删”；因此它适合工程链路，而不是一次性压测脚本。"
    )

    add_heading(doc, "六、三条主线配置到底谁好谁差", level=1)
    add_table(
        doc,
        headers=["配置", "Stage2", "Stage3", "有效总压缩率", "结论"],
        rows=[
            ["legacy_s12 + exact_only", "stage2_aggressive / linewise", "A only", "11.059100%", "老基线，可参考但非最优"],
            ["aggressive_s12 + exact_only", "stage2_hybrid_ab_aggressive / blockwise", "A only", "14.893045%", "推荐默认主线"],
            ["aggressive_s12 + hybrid", "stage2_hybrid_ab_aggressive / blockwise", "A+B", "14.906111%", "纯数字最优，但只高 0.013066pp"],
        ],
    )
    add_bullets(
        doc,
        [
            "为什么默认推荐 exact_only：只比 hybrid 低 0.013066pp，但复杂度和行为风险更低。",
            "为什么不继续用 legacy_s12：相对 aggressive_s12+exact_only 少了 3.833945pp，差距太大。",
            "一句话：要工程稳定就 exact_only，要刷榜极限才考虑 hybrid。"
        ],
    )

    add_heading(doc, "七、远端 GPU 两组训练对照（这次最关键）", level=1)
    add_table(
        doc,
        headers=["组别", "训练配方", "n_train_rows", "train_loss", "raw pass@1", "comp pass@1", "判定"],
        rows=[
            ["组A", "CPT(77) -> SFT(154)", "154", "0.8589", "0.0", "0.0", "不可用"],
            ["组B（最佳）", "CPT(77) -> SFT(220)", "220", "0.8440", "0.5", "0.4", "可用，且有真实 token 降低"],
        ],
    )
    add_bullets(
        doc,
        [
            "同样是 CPT(77) 起步，SFT 从 154 扩到 220 后，结果从 0.0/0.0 回到 0.5/0.4。",
            "说明这阶段数据规模与多样性比“严格配置对齐”更重要。",
            "也说明模型并非学不会，而是需要足够训练覆盖才能进入可用区间。",
        ],
    )

    add_heading(doc, "八、最佳组端到端指标拆解（为什么说真正省了）", level=1)
    add_table(
        doc,
        headers=["指标", "raw", "compressed", "delta", "解释"],
        rows=[
            ["avg_context_tokens", "1118.4", "979.4", "-139.0 (-12.4%)", "上下文已明显变短"],
            ["avg_prompt_tokens", "1246.2", "1110.1", "-136.1 (-10.9%)", "生成时实际提示也变短"],
            ["avg_codebook_tokens", "-", "0.6", "开销很小", "说明主要收益未被 codebook 吃掉"],
            ["pass@1", "0.5", "0.4", "-0.1", "能力尚未持平，仍需优化"],
        ],
    )
    add_paragraph(
        doc,
        "这组数据的意义是：我们第一次同时看到“真实提示变短”和“仍有可用准确率”，这和早期仅静态压缩好看是本质不同的里程碑。"
    )

    add_heading(doc, "九、quick eval 怎么读（很多人最容易误解的部分）", level=1)
    add_table(
        doc,
        headers=["组别", "expansion_success_rate", "leakage_rate", "compression_hit_rate", "解释"],
        rows=[
            ["CPT77+SFT154", "100%", "1350%", "0%", "能展开但泄漏极高，规则遵循很差"],
            ["CPT77+SFT220", "100%", "350%", "0%", "泄漏明显下降，但仍未学会稳定压缩输出"],
        ],
    )
    add_bullets(
        doc,
        [
            "expansion_success_rate=100% 不代表模型“全都学会”，它只说明展开流程能跑。",
            "真正关键的是 leakage_rate 和 compression_hit_rate，这两项仍然偏差。",
            "这也解释了为什么 HumanEval 上 compressed 仍掉点：规则控制还不够稳。",
        ],
    )

    add_heading(doc, "十、好在哪，差在哪（你最关心的判断）", level=1)
    add_heading(doc, "10.1 目前做得好的地方", level=2)
    add_bullets(
        doc,
        [
            "端到端 token 真实下降已被证明（不是只看静态压缩）。",
            "训练-评测链路完整：CPT -> SFT -> quick eval -> HumanEval 可闭环运行。",
            "规则系统可解释：每个“压/不压”基本都能回溯到明确规则。",
            "大 SFT 数据集显著改善可用性，说明方向有效而非随机撞运气。",
        ],
    )
    add_heading(doc, "10.2 当前仍然差的地方", level=2)
    add_bullets(
        doc,
        [
            "准确率未持平：compressed 仍比 raw 低 0.1。",
            "规则切换控制不稳：leakage_rate 高，compression_hit_rate 仍为 0%。",
            "B 通道在当前配置下贡献有限，复杂度收益比不高。",
            "部分配置和指标命名对非研发读者不友好，需要配套解释文档（本报告已补）。",
        ],
    )

    add_heading(doc, "十一、为什么现在不能说失败，也不能说大获全胜", level=1)
    add_paragraph(
        doc,
        "如果只看准确率，确实还有差距；但如果只看 token，又已经非常漂亮。工程上必须两者同时看。最客观的阶段判断是："
    )
    add_bullets(
        doc,
        [
            "已经跨过“压缩收益被包装开销吃掉”的阶段；",
            "尚未跨过“更短且不掉点”的最终目标线；",
            "当前状态可定义为“里程碑达成，但尚未产品化达标”。",
        ],
    )

    add_heading(doc, "十二、下一步怎么做才最有性价比", level=1)
    add_numbered(
        doc,
        [
            "先攻规则遵循：把 leakage_rate 继续压低，并把 compression_hit_rate 拉起来。",
            "继续扩充高质量 SFT 数据，优先覆盖容易触发别名展开/保留错误的样本。",
            "针对 -0.1 掉点做题级误差分析，区分规则失控与模型推理失误。",
            "保持 exact_only 作为默认主线，hybrid 仅做对照分支，不要过早切默认。",
            "每轮都固定产出“可读解释版报告”，避免只给数字不讲机制。",
        ],
    )

    add_heading(doc, "附录 A：本版使用的关键结果文件", level=1)
    add_bullets(
        doc,
        [
            "results/langv1_gpu_run/checkpoints/qwen15b_eval80_cpt_sft/train_summary.json",
            "results/langv1_gpu_run/checkpoints/qwen15b_eval80_cpt_sft128/train_summary.json",
            "results/langv1_gpu_run/humaneval/summary.json",
            "results/langv1_gpu_run/humaneval_sft128/summary.json",
            "results/langv1_gpu_run/evaluate_langv1_eval80.txt",
            "results/langv1_gpu_run/evaluate_langv1_sft128.txt",
            "docs/BEST_COMPRESSION_REMEASURE_2026-04-17.md",
            "docs/LANGV1_QWEN15B_RESULTS.md",
        ],
    )


def main() -> int:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _set_doc_defaults(doc)
    _build_report(doc)
    save_path = OUT_PATH
    try:
        doc.save(save_path)
    except PermissionError:
        save_path = DOCS_DIR / "etv2_detailed_conclusion_report_2026-04-19_formal_v2.docx"
        doc.save(save_path)
    print(save_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
