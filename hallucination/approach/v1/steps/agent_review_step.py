
import json
import re
import sys
import shutil
from datetime import datetime
from typing import List, Dict, Any, Optional
from openai.types.chat.chat_completion_message_param import ChatCompletionMessageParam
from pathlib import Path

pp_dir = str(Path(__file__).parent.parent)
if pp_dir not in sys.path:
    sys.path.append(str(pp_dir))

from utils.Generator import Generator
from utils.api_key import api_keys
from path_config import cfg_review_dir_path, cfg_review_prompt_file_path, cfg_translate_result_dir_path, cfg_binary_graph_file_path
from graph.retrieval_tools import sort_headers, load_nodes
from graph.structure import Header
from agent.Agent import CppCompilationAgent

# 尝试导入openpyxl，如果不存在则设置标志
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    print("警告: openpyxl未安装，将无法生成Excel统计文件。可使用 pip install openpyxl 安装。")


class CompilationStats:
    """编译统计数据类"""
    def __init__(self):
        self.records: List[Dict[str, Any]] = []
    
    def add_record(self, file_name: str, compile_attempts: int, status: str, error_msg: str = ""):
        """添加一条编译记录"""
        self.records.append({
            "file_name": file_name,
            "compile_attempts": compile_attempts,
            "status": status,
            "error_msg": error_msg
        })
    
    def get_summary(self) -> Dict[str, Any]:
        """获取统计摘要"""
        if not self.records:
            return {}
        
        total_files = len(self.records)
        total_attempts = sum(r["compile_attempts"] for r in self.records)
        success_count = sum(1 for r in self.records if r["status"] == "success")
        one_shot_count = sum(1 for r in self.records if r["compile_attempts"] == 1 and r["status"] == "success")
        max_attempts = max(r["compile_attempts"] for r in self.records) if self.records else 0
        avg_attempts = total_attempts / total_files if total_files > 0 else 0
        
        return {
            "total_files": total_files,
            "total_attempts": total_attempts,
            "success_count": success_count,
            "fail_count": total_files - success_count,
            "one_shot_count": one_shot_count,
            "max_attempts": max_attempts,
            "avg_attempts": round(avg_attempts, 2)
        }
    
    def export_to_excel(self, output_path: Path) -> bool:
        """导出统计数据到Excel文件"""
        if not OPENPYXL_AVAILABLE:
            print("错误: openpyxl未安装，无法生成Excel文件")
            return False
        
        if not self.records:
            print("警告: 没有编译记录，跳过Excel生成")
            return False
        
        try:
            wb = Workbook()
            sheet = wb.active
            sheet.title = "编译统计"
            
            # 样式定义
            header_fill = PatternFill('solid', fgColor='4472C4')
            header_font = Font(bold=True, color='FFFFFF')
            success_font = Font(color='008000')
            fail_font = Font(color='FF0000')
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # 表头
            headers = ["序号", "文件名", "编译次数", "状态", "备注"]
            for col, header in enumerate(headers, 1):
                cell = sheet.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center')
                cell.border = thin_border
            
            # 数据行
            for idx, record in enumerate(self.records, 1):
                row = idx + 1
                
                # 序号
                cell = sheet.cell(row=row, column=1, value=idx)
                cell.alignment = Alignment(horizontal='center')
                cell.border = thin_border
                
                # 文件名
                cell = sheet.cell(row=row, column=2, value=record["file_name"])
                cell.border = thin_border
                
                # 编译次数
                cell = sheet.cell(row=row, column=3, value=record["compile_attempts"])
                cell.alignment = Alignment(horizontal='center')
                cell.border = thin_border
                
                # 状态
                status_text = "成功" if record["status"] == "success" else "失败"
                cell = sheet.cell(row=row, column=4, value=status_text)
                cell.alignment = Alignment(horizontal='center')
                cell.border = thin_border
                if record["status"] == "success":
                    cell.font = success_font
                else:
                    cell.font = fail_font
                
                # 备注
                cell = sheet.cell(row=row, column=5, value=record.get("error_msg", ""))
                cell.border = thin_border
            
            # 统计汇总区域
            summary = self.get_summary()
            summary_start = len(self.records) + 4
            
            title_cell = sheet.cell(row=summary_start, column=1, value="统计汇总")
            title_cell.font = Font(bold=True, size=12)
            
            summary_items = [
                ("总文件数", summary["total_files"]),
                ("成功文件数", summary["success_count"]),
                ("失败文件数", summary["fail_count"]),
                ("总编译次数", summary["total_attempts"]),
                ("平均编译次数", summary["avg_attempts"]),
                ("一次成功数", summary["one_shot_count"]),
                ("最多编译次数", summary["max_attempts"]),
            ]
            
            for i, (label, value) in enumerate(summary_items):
                row = summary_start + 1 + i
                label_cell = sheet.cell(row=row, column=1, value=label)
                label_cell.border = thin_border
                
                value_cell = sheet.cell(row=row, column=2, value=value)
                value_cell.border = thin_border
                value_cell.alignment = Alignment(horizontal='center')
            
            # 设置列宽
            sheet.column_dimensions['A'].width = 12
            sheet.column_dimensions['B'].width = 30
            sheet.column_dimensions['C'].width = 12
            sheet.column_dimensions['D'].width = 10
            sheet.column_dimensions['E'].width = 40
            
            # 保存文件
            wb.save(output_path)
            print(f"编译统计Excel已保存到: {output_path}")
            return True
            
        except Exception as e:
            print(f"生成Excel文件时出错: {e}")
            return False
    
    def export_to_csv(self, output_path: Path) -> bool:
        """导出统计数据到CSV文件（备用方案，不需要openpyxl）"""
        if not self.records:
            print("警告: 没有编译记录，跳过CSV生成")
            return False
        
        try:
            import csv
            with open(output_path, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.writer(f)
                
                # 表头
                writer.writerow(["序号", "文件名", "编译次数", "状态", "备注"])
                
                # 数据行
                for idx, record in enumerate(self.records, 1):
                    status_text = "成功" if record["status"] == "success" else "失败"
                    writer.writerow([
                        idx,
                        record["file_name"],
                        record["compile_attempts"],
                        status_text,
                        record.get("error_msg", "")
                    ])
                
                # 空行
                writer.writerow([])
                
                # 统计汇总
                summary = self.get_summary()
                writer.writerow(["统计汇总"])
                writer.writerow(["总文件数", summary["total_files"]])
                writer.writerow(["成功文件数", summary["success_count"]])
                writer.writerow(["失败文件数", summary["fail_count"]])
                writer.writerow(["总编译次数", summary["total_attempts"]])
                writer.writerow(["平均编译次数", summary["avg_attempts"]])
                writer.writerow(["一次成功数", summary["one_shot_count"]])
                writer.writerow(["最多编译次数", summary["max_attempts"]])
            
            print(f"编译统计CSV已保存到: {output_path}")
            return True
            
        except Exception as e:
            print(f"生成CSV文件时出错: {e}")
            return False


class AgentReviewStep:
    def __init__(self):
        self.stats = CompilationStats()

    def copy_files(self, ai_name, project_name):
        result_dir_path = cfg_translate_result_dir_path(ai_name, project_name)
        review_dir_path = cfg_review_dir_path(ai_name, project_name)
        if review_dir_path.exists():
            shutil.rmtree(review_dir_path)
        #复制文件
        shutil.copytree(result_dir_path, review_dir_path)

    def run_review(self, ai_name, project_name, generate_stats: bool = True):
        """
        运行审查流程
        
        Args:
            ai_name: AI模型名称
            project_name: 项目名称
            generate_stats: 是否生成编译统计文件，默认为True
        """
        print(f"Running review for {ai_name} on {project_name}")
        self.copy_files(ai_name, project_name)
        
        # 重置统计数据
        self.stats = CompilationStats()
        
        data = load_nodes(cfg_binary_graph_file_path(ai_name, project_name))
        sorted_headers = sort_headers(data["headers"]) 
        if sorted_headers is None:
            print("警告：由于存在循环依赖，无法确定理想的审查顺序。")
            # 方案：直接使用原始的、未排序的列表，避免程序终止
            all_headers = data["headers"] 
            if isinstance(all_headers, dict):
                # 如果是字典（key 为 header.key），转为列表
                process_list = list(all_headers.values())
            else:
                process_list = all_headers
            print("程序将按原始加载顺序继续处理。")
        else:
            process_list = sorted_headers
        print(f"review order: {[header.key for header in process_list]}") #type: ignore

        for header in process_list: #type: ignore
            file_path = cfg_review_dir_path(ai_name, project_name) / (header.key + ".cpp")
            if not file_path.exists():
                print(f"file skipped: {file_path}. Maybe an interface")
                continue

            api_key = api_keys.get(ai_name, "")
            # Create agent
            agent = CppCompilationAgent(ai_name, api_key, file_path)

            # Run agent
            result = agent.run(max_attempts=10)

            # 记录编译统计
            error_msg = result.get('error', result.get('message', ''))
            self.stats.add_record(
                file_name=header.key + ".cpp",
                compile_attempts=result['attempts'],
                status=result['status'],
                error_msg=error_msg if result['status'] != 'success' else ''
            )

            # Print results
            print("\n" + "="*50)
            print("Agent Execution Complete")
            print("="*50)
            print(f"Status: {result['status']}")
            print(f"Attempts: {result['attempts']}")
            if 'error' in result:
                print(f"Error: {result['error']}")
            if 'final_response' in result:
                print(f"Final Response: {result['final_response']}")
            print("\n" + "="*50)

        # 生成统计文件
        if generate_stats:
            self._generate_stats_file(ai_name, project_name)
    
    def _generate_stats_file(self, ai_name: str, project_name: str):
        """生成编译统计文件"""
        review_dir = cfg_review_dir_path(ai_name, project_name)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        
        # 优先生成Excel，如果openpyxl不可用则生成CSV
        if OPENPYXL_AVAILABLE:
            output_path = review_dir / f"compilation_stats_{timestamp}.xlsx"
            success = self.stats.export_to_excel(output_path)
        else:
            output_path = review_dir / f"compilation_stats_{timestamp}.csv"
            success = self.stats.export_to_csv(output_path)
        
        if success:
            # 打印统计摘要
            summary = self.stats.get_summary()
            print("\n" + "="*50)
            print("编译统计摘要")
            print("="*50)
            print(f"总文件数: {summary['total_files']}")
            print(f"成功: {summary['success_count']}, 失败: {summary['fail_count']}")
            print(f"总编译次数: {summary['total_attempts']}")
            print(f"平均编译次数: {summary['avg_attempts']}")
            print(f"一次成功数: {summary['one_shot_count']}")
            print(f"最多编译次数: {summary['max_attempts']}")
            print("="*50 + "\n")
    
    def get_stats(self) -> CompilationStats:
        """获取编译统计对象"""
        return self.stats


