# import os
# import re
# from openpyxl import Workbook, load_workbook
# from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


# def find_file(dir_path:str, file_name:str)->str:
#     '''
#     查找文件
#     '''
#     for root, dirs_names, file_names in os.walk(dir_path):
#         if file_name in file_names:
#             return os.path.join(root, file_name)
#     return None


# def get_comment(code:str, method_name:str)->tuple[int, str]:
#     '''
#     获取方法的注释
#     '''  
#     try:
#         comment_pattern = re.compile(r'.*{}.*//#(\d*)(.*)'.format(method_name))
#     except re.PatternError:
#         return -1,f'error: comment pattern error, method name is {method_name}'
#     match = re.search(comment_pattern, code)
#     if match:
#         return int(match.group(1) or 0), match.group(2)
#     else:
#         if (method_name + '(') in code:
#             return -1, 'error: no comment found'
#         else:
#             return 4, '方法未实现'

# def over_all_comment(code:str)->str:
#     '''
#     获取整体注释
#     '''
#     comment_pattern = re.compile(r'//##(\d*)(.*)')
#     match = re.search(comment_pattern, code)
#     if match:
#         return int(match.group(1) or 0), match.group(2)
#     else:
#         return 0, 'error: no overall comment found'
    
# def get_method_name(split_content:str)->list[str]:
#     '''
#     获取方法名
#     '''
#     method_names = []
#     pattern = r'```method(.*?)```'
#     matches = re.findall(pattern, split_content, re.DOTALL)
#     for match in matches:
#         names = re.findall(r'//name:(.*)\n', match)
#         if len(names) != 1:
#             print(f'warning: find{len(names)} method names in {match}')
#         method_names.append(names[0])
#     return method_names



# def deal_single_file(split_file_path, translated_file_path1, translated_file_path2, row_num, worksheet):
#     worksheet['B' + str(row_num)] = os.path.basename(split_file_path)
#     print(f"deal with {translated_file_path1}")
#     if not os.path.exists(translated_file_path1):
        
#         worksheet['C' + str(row_num)] = 'error: translated file not found'
#         worksheet['D' + str(row_num)] = 'error: translated file not found'
#         worksheet['E' + str(row_num)] = -1
#         # raise FileNotFoundError(f'{translated_file_path} not found')
#         return row_num + 1

#     try:
#         with open(split_file_path, 'r', encoding='utf-8') as f:
#             split_content = f.read()
#     except UnicodeDecodeError:
#         worksheet['C' + str(row_num)] = 'error: source file encoding error'
#         worksheet['D' + str(row_num)] = 'error: source file encoding error'
#         worksheet['E' + str(row_num)] = -1
#         return row_num + 1
    
#     try:
#         translated_code1 = None
#         translated_code2 = None
#         with open (translated_file_path1, 'r', encoding='utf-8') as f:
#             translated_code1 = f.read()
#         if translated_file_path2 is not None:
#             with open (translated_file_path2, 'r', encoding='utf-8') as f:
#                 translated_code2 = f.read()
#     except UnicodeDecodeError:
#         worksheet['C' + str(row_num)] = 'error: translated file encoding error'
#         worksheet['D' + str(row_num)] = 'error: translated file encoding error'
#         worksheet['E' + str(row_num)] = -1
#         return row_num + 1
    

#     methods = get_method_name(split_content)
#     print(methods)

#     index, overall_comment = over_all_comment(translated_code1)
#     if overall_comment == 'error: no overall comment found' and translated_code2 is not None:
#         index, overall_comment = over_all_comment(translated_code2)

#     worksheet['D' + str(row_num)] = overall_comment
#     start_row = row_num

#     for method in methods:
#         worksheet['C' + str(row_num)] = method
#         print(f"C{row_num}写入{method}")

#         index, comment = get_comment(translated_code1, method)
#         if index == -1 and translated_code2 is not None:
#             index, comment = get_comment(translated_code2, method)
#         worksheet['E' + str(row_num)] = index
#         worksheet['F' + str(row_num)] = comment
#         row_num += 1

#     if row_num == start_row:
#         worksheet['C' + str(row_num)] = 'error: no method found'
#         worksheet['D' + str(row_num)] = 'error: no method found'
#         worksheet['E' + str(row_num)] = -1
#         row_num += 1

#     worksheet.merge_cells('B' + str(start_row) + ':B' + str(row_num - 1))
#     worksheet.merge_cells('D' + str(start_row) + ':D' + str(row_num - 1))

#     return row_num

    
# def deal_dir(dir_path, row_num, worksheet, AI_name, strategy):
#     dir_name = os.path.basename(dir_path)
#     worksheet['A' + str(row_num)] = dir_name
#     start_row = row_num

#     print(f'load data from {os.path.join(dir_path, AI_name, strategy)}...')
#     if not os.path.exists(os.path.join(dir_path, AI_name, strategy)):
#         print(f'warning: {os.path.join(dir_path, AI_name, strategy)} not found')
#         worksheet['B' + str(row_num)] = 'error: translated dir not found'
#         worksheet['C' + str(row_num)] = 'error: translated dir not found'
#         worksheet['D' + str(row_num)] = 'error: translated dir not found'
#         return row_num + 1

#     for file_name in os.listdir(os.path.join(dir_path, 'split')):
#         split_file_path = os.path.join(dir_path,'split', file_name)
#         translated_file_name1 = file_name.replace('.txt', '.h')
#         translated_file_name2 = file_name.replace('.txt', '.cpp')
#         translated_file_path1 = os.path.join(dir_path, AI_name, strategy, translated_file_name1)
#         translated_file_path2 = os.path.join(dir_path, AI_name, strategy, translated_file_name2) if os.path.exists(os.path.join(dir_path, AI_name, strategy, translated_file_name2)) else None
        
#         row_num = deal_single_file(split_file_path, translated_file_path1, translated_file_path2, row_num, worksheet)

#     worksheet.merge_cells('A' + str(start_row) + ':A' + str(row_num - 1))

#     return row_num

# def create_data_set(dir_path, data_set_path, AI_name, data_set_sources, strategy):
#     if os.path.exists(data_set_path):
#         os.remove(data_set_path)
        
    
#     workbook = Workbook()
#     worksheet = workbook.active
#     worksheet.title = 'data_set'
#     row_num = 1

#     for project_name in data_set_sources:
#         sub_dir_path = os.path.join(dir_path, project_name)
#         if os.path.isdir(sub_dir_path):
#             row_num = deal_dir(sub_dir_path, row_num, worksheet, AI_name, strategy)
#         else:
#             print(f"数据集源目录{sub_dir_path}不存在，请检查")

#     workbook.save(data_set_path)

    








import os
import re
import csv
import pickle
from typing import List, Tuple, Optional, Dict, Any

# 错误码列（左到右）
ERROR_CODES: List[int] = [-1, 0] + list(range(1, 12))  # -1,0,1..11
ERROR_COL_NAMES: List[str] = [f"Err_{c}" for c in ERROR_CODES]

# =========================
# 工具：兼容老版本 pickle 的模块别名
# =========================

def _install_pickle_aliases():
    """
    老的 nodes.pkl 里，类模块可能记录成顶层 'structure' 等。
    在反序列化前把这些老名字映射到当前能导入到的模块，避免报错。
    """
    import sys, importlib
    alias_map = {
        "structure": [
            "sitp_script.strategy_bottomup.structure",
            "sitp_script.structure",
            "structure",
        ],
        "retrieval_tools": [
            "sitp_script.strategy_bottomup.retrieval_tools",
            "sitp_script.retrieval_tools",
            "retrieval_tools",
        ],
        "Generator": [
            "sitp_script.Generator",
            "Generator",
        ],
        "api_key": [
            "sitp_script.api_key",
            "api_key",
        ],
    }
    for alias, candidates in alias_map.items():
        if alias in sys.modules:
            continue
        for cand in candidates:
            try:
                mod = importlib.import_module(cand)
                sys.modules[alias] = mod
                break
            except Exception:
                continue


# =========================
# 解析与规范化
# =========================

def _parse_error_segments(seg_str: str) -> List[Tuple[int, str]]:
    """
    '3Missing!8Pointer' -> [(3,'Missing'), (8,'Pointer')]
    'None' / '#None'    -> [(0,'None')]
    其他不符合 -> [(-1,'invalid tag: xxx')]（注意：这个 -1 仅用于表示“标注格式非法”，
    最终仍会走规范化逻辑保障 0/-1 互斥）
    """
    results: List[Tuple[int, str]] = []
    if not seg_str:
        return results
    for raw in [s.strip() for s in seg_str.split('!') if s.strip()]:
        m = re.match(r'(\-?\d+)\s*(.*)', raw)
        if m:
            code = int(m.group(1))
            text = (m.group(2) or '').strip()
            results.append((code, text))
        else:
            if raw.lower() in ('none', '#none'):
                results.append((0, 'None'))
            else:
                results.append((-1, f'invalid tag: {raw}'))
    return results


def _normalize_segments(segments: List[Tuple[int, str]]) -> List[Tuple[int, str]]:
    """
    执行『0 与 -1 互斥』：
      - 若含 0（None）=> 仅保留 [(0,'None')]
      - 若含 -1      => 仅保留第一条 (-1, msg)
      - 否则保持原列表（支持多错误）
    """
    if not segments:
        return []
    for c, _ in segments:
        if c == 0:
            return [(0, 'None')]
    for c, t in segments:
        if c == -1:
            return [(-1, t)]
    return segments


# =========================
# 读取/写回 nodes.pkl
# =========================

def _nodes_pkl_path(dir_path: str, project: str) -> str:
    return os.path.join(dir_path, "output_info", project, "nodes.pkl")

def _load_nodes(dir_path: str, project: str) -> Optional[Dict[str, Any]]:
    p = _nodes_pkl_path(dir_path, project)
    if not os.path.exists(p):
        return None
    # 安装别名后再反序列化
    import importlib, sys  # noqa
    _install_pickle_aliases()
    with open(p, "rb") as f:
        return pickle.load(f)

def _save_nodes(dir_path: str, project: str, headers, method_nodes) -> None:
    p = _nodes_pkl_path(dir_path, project)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "wb") as f:
        pickle.dump({"headers": headers, "method_nodes": method_nodes}, f)

def _build_index(headers, method_layers):
    """
    返回：
      header_by_key:            {header.key -> Header}
      node_by_hdr_and_name:     {(header.key, method_name) -> Node}
    """
    header_by_key: Dict[str, Any] = {}
    for h in headers or []:
        k = getattr(h, "key", None)
        if k:
            header_by_key[k] = h

    node_by_hdr_and_name: Dict[Tuple[str, str], Any] = {}

    # 优先从 header.methods 建
    for h in headers or []:
        k = getattr(h, "key", None)
        if not k:
            continue
        for n in getattr(h, "methods", []) or []:
            node_key = getattr(n, "key", "")
            short_name = node_key.split(":", 1)[-1] if ":" in node_key else node_key
            if short_name:
                node_by_hdr_and_name.setdefault((k, short_name), n)

    # 兜底：再扫所有层
    for layer in method_layers or []:
        for n in layer:
            header = getattr(n, "header", None)
            hk = getattr(header, "key", None)
            node_key = getattr(n, "key", "")
            short_name = node_key.split(":", 1)[-1] if ":" in node_key else node_key
            if hk and short_name and (hk, short_name) not in node_by_hdr_and_name:
                node_by_hdr_and_name[(hk, short_name)] = n

    return header_by_key, node_by_hdr_and_name


# =========================
# .h/.cpp 解析：总体与方法
# =========================

# 总体注释（单文件）
def _overall_from_code(code: Optional[str]) -> List[Tuple[int, str]]:
    if not code:
        return [(-1, 'error: no overall comment found')]
    m = re.search(r'//\#\#([^\n]+)', code)
    if not m:
        return [(-1, 'error: no overall comment found')]
    segs = _parse_error_segments(m.group(1).strip())
    segs = _normalize_segments(segs)
    return segs if segs else [(-1, 'error: empty overall comment')]

# .h 中方法声明（宽松）
_H_METHOD_DECL_RE = re.compile(
    r'^[ \t]*(?:template\s*<[^>]+>\s*)?'
    r'(?:virtual|static|inline|constexpr|explicit|friend)?[ \t]*'
    r'[\w:\<\>\*\&\s~]+?\b(~?[A-Za-z_]\w*)\s*\([^;{]*\)\s*'
    r'(?:const\s*)?(?:=\s*0\s*)?;',
    re.MULTILINE
)

# .cpp 中方法定义（Class::Method(...)）
_CPP_METHOD_DEF_RE = re.compile(
    r'^[ \t]*(?:template\s*<[^>]+>\s*)?'
    r'[\w:\<\>\*\&\s~]*?\b(?:[A-Za-z_]\w*::)+(~?[A-Za-z_]\w*)\s*\(',
    re.MULTILINE
)

_OPERATOR_NAME_RE = re.compile(r'\boperator\s*[^\s(]+\b')

def _extract_methods_from_h(code: Optional[str]) -> List[str]:
    if not code:
        return []
    names: List[str] = []
    for m in _H_METHOD_DECL_RE.finditer(code):
        nm = m.group(1)
        if nm and not _OPERATOR_NAME_RE.search(nm):
            names.append(nm)
    return names

def _extract_methods_from_cpp(code: Optional[str]) -> List[str]:
    if not code:
        return []
    names: List[str] = []
    for m in _CPP_METHOD_DEF_RE.finditer(code):
        nm = m.group(1)
        if nm and not _OPERATOR_NAME_RE.search(nm):
            names.append(nm)
    return names

def _get_inline_comment_for_method(code: Optional[str], method_name: str) -> Optional[List[Tuple[int, str]]]:
    """匹配“所在行” //#... 标注"""
    if not code:
        return None
    pattern = re.compile(rf'^[^\n]*\b{re.escape(method_name)}\b[^\n]*//\#([^\n]+)$', re.MULTILINE)
    m = pattern.search(code)
    if not m:
        return None
    segs = _parse_error_segments(m.group(1).strip())
    segs = _normalize_segments(segs)
    return segs if segs else [(-1, 'error: empty comment')]


# =========================
# 组装单文件：返回多行 & 回写 nodes
# =========================

def _segments_to_cols(segs: List[Tuple[int, str]]) -> Dict[int, str]:
    """
    把 [(code, text), ...] 汇总到列字典：{code -> 'text1 | text2 ...'}
    """
    bucket: Dict[int, List[str]] = {}
    for code, text in segs:
        bucket.setdefault(code, []).append(text or '')
    return {code: ' | '.join(v) for code, v in bucket.items()}

def _collect_rows_for_file_and_update_graph(
    dir_path: str,
    project: str,
    file_stem: str,
    h_path: Optional[str],
    cpp_path: Optional[str],
    header_by_key: Dict[str, Any],
    node_by_hdr_and_name: Dict[Tuple[str, str], Any],
) -> List[List[object]]:
    rows: List[List[object]] = []
    file_base = f"{file_stem}.txt"

    # 读 .h / .cpp
    def _read(path):
        if path and os.path.exists(path):
            try:
                with open(path, 'r', encoding='utf-8') as f:
                    return f.read()
            except UnicodeDecodeError:
                return None
        return None

    h_code = _read(h_path)
    cpp_code = _read(cpp_path)

    # ===== 总体：先 .h 再 .cpp；无则 -1 =====
    overall = _overall_from_code(h_code)
    if overall == [(-1, 'error: no overall comment found')]:
        overall = _overall_from_code(cpp_code)

    # CSV（总体一行，Method 留空；错误填入对应列）
    overall_cols = _segments_to_cols(overall)
    row = [project, file_base, ""]
    for code in ERROR_CODES:
        row.append(overall_cols.get(code, ""))
    rows.append(row)

    # 回写 Header.overall_errors
    hdr = header_by_key.get(file_stem)
    if hdr is not None:
        try:
            setattr(hdr, "overall_errors", overall)
        except Exception:
            pass

    # ===== 方法名：合并 .h/.cpp 去重 =====
    method_names: List[str] = []
    seen = set()
    for nm in _extract_methods_from_h(h_code) + _extract_methods_from_cpp(cpp_code):
        if nm not in seen:
            seen.add(nm)
            method_names.append(nm)

    if not method_names:
        # 仍写一行提示“无方法”，放在 Err_-1 列
        row = [project, file_base, "error: no method found"]
        for code in ERROR_CODES:
            row.append("error: no method found" if code == -1 else "")
        rows.append(row)
        return rows

    # ===== 每个方法：优先 .h 再 .cpp；都无 => -1 =====
    for method in method_names:
        segs = _get_inline_comment_for_method(h_code, method)
        if segs is None:
            segs = _get_inline_comment_for_method(cpp_code, method)
        if segs is None or len(segs) == 0:
            segs = [(-1, 'error: no comment found')]

        cols = _segments_to_cols(segs)
        row = [project, file_base, method]
        for code in ERROR_CODES:
            row.append(cols.get(code, ""))
        rows.append(row)

        # 回写 Node.errors
        node = None
        if hdr is not None:
            node = node_by_hdr_and_name.get((getattr(hdr, "key", file_stem), method))
        if node is None:
            node = node_by_hdr_and_name.get((file_stem, method))
        if node is not None:
            try:
                setattr(node, "errors", segs)
            except Exception:
                pass

    return rows


# =========================
# 主入口：创建 CSV/XLSX 并回写 nodes.pkl
# =========================

def create_data_set(dir_path: str, data_set_path: str, AI_name: str, data_set_sources: List[str], strategy: str, save_csv: bool = True):
    """
    dir_path:          translation_java-cpp 根目录
    data_set_path:     输出文件（建议放 README 同级）。若以 .xlsx 结尾，也会尝试写 xlsx；但无论如何都会写同名 .csv
    AI_name:           'deepseek' / 'gpt'
    data_set_sources:  项目名列表（dir_path 下子目录）
    strategy:          'class' | 'method' | 'bottomup'（用于定位 <AI_name>/<strategy> 下的 .h/.cpp）
    """
    # 先准备所有行，包含表头
    rows: List[List[object]] = []
    header_row = ["Project", "File", "Method"] + ERROR_COL_NAMES
    rows.append(header_row)

    for project_name in data_set_sources:
        sub_dir = os.path.join(dir_path, project_name)
        split_dir = os.path.join(sub_dir, 'split')
        trans_dir = os.path.join(sub_dir, AI_name, strategy)

        nodes_data = _load_nodes(dir_path, project_name)
        if nodes_data:
            headers = nodes_data.get("headers", [])
            method_layers = nodes_data.get("method_nodes", [])
            header_by_key, node_by_hdr_and_name = _build_index(headers, method_layers)
        else:
            headers, method_layers = None, None
            header_by_key, node_by_hdr_and_name = {}, {}

        if not os.path.isdir(split_dir) or not os.path.isdir(trans_dir):
            # 目录缺失也写一行错误提示
            row = [project_name, 'error: translated dir not found', ""]
            for code in ERROR_CODES:
                row.append("error: translated dir not found" if code == -1 else "")
            rows.append(row)
            if nodes_data:
                _save_nodes(dir_path, project_name, headers, method_layers)
            continue

        # 遍历 split 下的 .txt 文件以获得 file_stem 集合
        for file_name in os.listdir(split_dir):
            if not file_name.endswith('.txt'):
                continue
            stem = os.path.splitext(file_name)[0]
            h_path   = os.path.join(trans_dir, stem + '.h')
            cpp_path = os.path.join(trans_dir, stem + '.cpp')

            rows.extend(
                _collect_rows_for_file_and_update_graph(
                    dir_path,
                    project_name,
                    stem,
                    h_path,
                    cpp_path,
                    header_by_key,
                    node_by_hdr_and_name,
                )
            )

        # 回写 nodes.pkl（若已加载）
        if nodes_data:
            _save_nodes(dir_path, project_name, headers, method_layers)

    # 写 CSV（一定写）
    if data_set_path.lower().endswith('.csv'):
        csv_path = data_set_path
    else:
        csv_path = os.path.splitext(data_set_path)[0] + '.csv'

    os.makedirs(os.path.dirname(csv_path) or '.', exist_ok=True)
    with open(csv_path, 'w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        for r in rows:
            writer.writerow(r)

    # 若目标为 .xlsx，则也写一份 xlsx（附带表头）
    if data_set_path.lower().endswith('.xlsx'):
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = 'data_set'
            for r in rows:
                ws.append(r)
            wb.save(data_set_path)
        except Exception as e:
            print(f'warning: failed to write xlsx: {e}')

    return csv_path


